"""
utils.py — Shared configuration, in-memory state, helper functions, and DB operations.
All cogs import from here so that state (active_sessions, etc.) is shared across modules.
"""

import os
import re
import time
import calendar
import asyncio
import csv
import io
from datetime import datetime, timedelta, time as dt_time, timezone
from zoneinfo import ZoneInfo

import discord
from discord.ext import commands
import aiosqlite
import matplotlib
matplotlib.use("Agg")  # Render graphs without a display (VPS-friendly)
import matplotlib.pyplot as plt
from dotenv import load_dotenv

# ==========================================
# CONFIGURATION — LOADED FROM .ENV
# ==========================================
load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN") or os.getenv("DISCORD_TOKEN")
if not BOT_TOKEN:
    raise ValueError("❌ BOT_TOKEN or DISCORD_TOKEN is missing in your .env file!")

DB_FILE = os.getenv("DB_FILE", "voice_stats.db")


def get_env_id(key: str) -> int | None:
    val = os.getenv(key)
    if not val:
        return None
    clean = val.split("#")[0].strip()
    digits = re.findall(r"\d+", clean)
    return int(digits[0]) if digits else None


def get_env_allowed_roles() -> dict:
    raw_val = (
        os.getenv("ROLESTATS_ALLOWED_ROLES")
        or os.getenv("ALLOWED_ROLE_IDS")
        or os.getenv("ROLESTATS_ROLES")
        or os.getenv("ALLOWED_ROLES")
        or ""
    )
    if not raw_val:
        return {"ids": set(), "names": set()}

    clean_val = raw_val.split("#")[0].strip()
    if not clean_val:
        return {"ids": set(), "names": set()}

    parts = [p.strip().strip("'\"") for p in clean_val.split(",") if p.strip()]

    role_ids: set[int] = set()
    role_names: set[str] = set()

    for p in parts:
        digits = re.findall(r"\d+", p)
        # Discord snowflake IDs are >= 15 digits
        if digits and len(digits[0]) >= 15:
            role_ids.add(int(digits[0]))
        else:
            clean_name = p.lstrip("@").strip().lower()
            if clean_name:
                role_names.add(clean_name)

    return {"ids": role_ids, "names": role_names}


STATS_CHANNEL_ID = get_env_id("STATS_CHANNEL_ID")
REPORT_CHANNEL_ID = get_env_id("REPORT_CHANNEL_ID")
AFK_CHANNEL_ID = get_env_id("AFK_CHANNEL_ID")
ROLESTATS_ALLOWED_ROLES = get_env_allowed_roles()

tz_name = os.getenv("TIMEZONE", "Asia/Dhaka")
try:
    LOCAL_TZ = ZoneInfo(tz_name)
except Exception:
    try:
        LOCAL_TZ = ZoneInfo("Asia/Dhaka")
    except Exception:
        LOCAL_TZ = timezone(timedelta(hours=6))


def can_use_rolestats(member: discord.Member) -> bool:
    """Returns True if the member is allowed to use /rolestats."""
    if not isinstance(member, discord.Member):
        return False
    if member.guild_permissions.administrator:
        return True
    allowed = ROLESTATS_ALLOWED_ROLES
    if allowed["ids"] or allowed["names"]:
        for r in member.roles:
            if r.id in allowed["ids"]:
                return True
            if r.name.lower() in allowed["names"]:
                return True
    return False


# ==========================================
# IN-MEMORY TRACKING (shared across cogs)
# ==========================================
active_sessions: dict[int, dict] = {}
last_connected_channels: dict[int, str] = {}  # {user_id: last_channel_name}
deafen_timestamps: dict[int, float] = {}       # {user_id: timestamp_deafened}
afk_moved_timestamps: dict[int, float] = {}   # {user_id: timestamp_moved_to_afk}


# ==========================================
# HELPER FUNCTIONS
# ==========================================
def format_duration(seconds: float) -> str:
    """Converts raw seconds into a readable string (e.g., 2h 15m 30s)."""
    if not seconds or seconds <= 0:
        return "0s"
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    parts = []
    if h > 0:
        parts.append(f"{h}h")
    if m > 0:
        parts.append(f"{m}m")
    if s > 0 or not parts:
        parts.append(f"{s}s")
    return " ".join(parts)


def get_cutoff_dates(cutoff_mode: str = "current", ref_date: datetime | None = None) -> tuple[str, str, str]:
    """
    Returns (start_date, end_date, label) for the 1st or 2nd cutoff of the month.
    1st Cutoff: 1st - 15th
    2nd Cutoff: 16th - Last Day of Month
    """
    if ref_date is None:
        ref_date = datetime.now(LOCAL_TZ)

    year = ref_date.year
    month = ref_date.month
    day = ref_date.day
    _, last_day = calendar.monthrange(year, month)

    if cutoff_mode == "current":
        if day <= 15:
            start_date = f"{year}-{month:02d}-01"
            end_date = f"{year}-{month:02d}-15"
            label = f"1st Cutoff (1-15 {calendar.month_abbr[month]} {year})"
        else:
            start_date = f"{year}-{month:02d}-16"
            end_date = f"{year}-{month:02d}-{last_day}"
            label = f"2nd Cutoff (16-{last_day} {calendar.month_abbr[month]} {year})"
        return start_date, end_date, label

    elif cutoff_mode == "previous":
        if day <= 15:
            if month == 1:
                prev_year = year - 1
                prev_month = 12
            else:
                prev_year = year
                prev_month = month - 1
            _, prev_last_day = calendar.monthrange(prev_year, prev_month)
            start_date = f"{prev_year}-{prev_month:02d}-16"
            end_date = f"{prev_year}-{prev_month:02d}-{prev_last_day}"
            label = f"2nd Cutoff (16-{prev_last_day} {calendar.month_abbr[prev_month]} {prev_year})"
        else:
            start_date = f"{year}-{month:02d}-01"
            end_date = f"{year}-{month:02d}-15"
            label = f"1st Cutoff (1-15 {calendar.month_abbr[month]} {year})"
        return start_date, end_date, label

    return f"{year}-{month:02d}-01", f"{year}-{month:02d}-{last_day}", f"{calendar.month_name[month]} {year}"


def get_specific_cutoff_range(year: int, month: int, part: int) -> tuple[str, str, str]:
    """Returns (start_date, end_date, label) for a specific year, month, and cutoff part (1 or 2)."""
    _, last_day = calendar.monthrange(year, month)
    if part == 1:
        return (
            f"{year}-{month:02d}-01",
            f"{year}-{month:02d}-15",
            f"1st Cutoff (1-15 {calendar.month_abbr[month]} {year})",
        )
    else:
        return (
            f"{year}-{month:02d}-16",
            f"{year}-{month:02d}-{last_day}",
            f"2nd Cutoff (16-{last_day} {calendar.month_abbr[month]} {year})",
        )


def get_date_range(time_filter: str) -> tuple[str, str]:
    """Returns a (start_date, end_date) tuple as YYYY-MM-DD strings."""
    now = datetime.now(LOCAL_TZ)
    today_str = now.strftime("%Y-%m-%d")

    if time_filter in ("current_cutoff", "cutoff_current"):
        s_date, e_date, _ = get_cutoff_dates("current", now)
        return s_date, e_date
    elif time_filter in ("previous_cutoff", "cutoff_previous"):
        s_date, e_date, _ = get_cutoff_dates("previous", now)
        return s_date, e_date
    elif time_filter == "today":
        return today_str, today_str
    elif time_filter == "this_week":
        start_of_week = now - timedelta(days=now.weekday())
        return start_of_week.strftime("%Y-%m-%d"), today_str
    elif time_filter == "this_month":
        start_of_month = now.replace(day=1)
        return start_of_month.strftime("%Y-%m-%d"), today_str
    elif time_filter == "last_month":
        first_day_this_month = now.replace(day=1)
        last_day_last_month = first_day_this_month - timedelta(days=1)
        first_day_last_month = last_day_last_month.replace(day=1)
        return first_day_last_month.strftime("%Y-%m-%d"), last_day_last_month.strftime("%Y-%m-%d")
    elif time_filter == "this_year":
        start_of_year = now.replace(month=1, day=1)
        return start_of_year.strftime("%Y-%m-%d"), today_str
    elif time_filter == "all_time":
        return "2000-01-01", "2099-12-31"

    return today_str, today_str


def parse_custom_date(date_str: str) -> str | None:
    """Parses a date string in YYYY-MM-DD format. Returns None on failure."""
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str.strip(), "%Y-%m-%d").strftime("%Y-%m-%d")
    except ValueError:
        return None


def determine_state(voice_state: discord.VoiceState) -> str:
    """Returns the audio state string based on a VoiceState object."""
    if voice_state.self_deaf or voice_state.deaf:
        return "deafened"
    elif voice_state.self_mute or voice_state.mute:
        return "muted"
    else:
        return "unmuted"


# ==========================================
# DATABASE OPERATIONS
# ==========================================
async def init_db():
    """Initialises the database, upgrading schema if needed."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("PRAGMA table_info(voice_activity)")
        columns = [row[1] for row in await cursor.fetchall()]

        if columns and "month_year" in columns and "record_date" not in columns:
            print("🚀 Upgrading database schema for daily tracking...")
            await db.execute("ALTER TABLE voice_activity RENAME TO voice_activity_old")
            await db.execute("""
                CREATE TABLE IF NOT EXISTS voice_activity (
                    user_id INTEGER,
                    user_name TEXT,
                    channel_id INTEGER,
                    channel_name TEXT,
                    record_date TEXT,
                    unmuted_seconds REAL DEFAULT 0,
                    muted_seconds REAL DEFAULT 0,
                    deafened_seconds REAL DEFAULT 0,
                    PRIMARY KEY (user_id, channel_id, record_date)
                )
            """)
            await db.execute("""
                INSERT INTO voice_activity (user_id, user_name, channel_id, channel_name, record_date, unmuted_seconds, muted_seconds, deafened_seconds)
                SELECT user_id, user_name, channel_id, channel_name, month_year || '-01', unmuted_seconds, muted_seconds, deafened_seconds
                FROM voice_activity_old
            """)
            await db.execute("DROP TABLE voice_activity_old")
            print("✅ Database successfully upgraded!")
        else:
            await db.execute("""
                CREATE TABLE IF NOT EXISTS voice_activity (
                    user_id INTEGER,
                    user_name TEXT,
                    channel_id INTEGER,
                    channel_name TEXT,
                    record_date TEXT,
                    unmuted_seconds REAL DEFAULT 0,
                    muted_seconds REAL DEFAULT 0,
                    deafened_seconds REAL DEFAULT 0,
                    PRIMARY KEY (user_id, channel_id, record_date)
                )
            """)

        await db.execute("DELETE FROM voice_activity WHERE channel_id IS NULL OR channel_name IS NULL")
        await db.execute("CREATE INDEX IF NOT EXISTS idx_user_date ON voice_activity(user_id, record_date)")
        await db.execute("CREATE INDEX IF NOT EXISTS idx_date ON voice_activity(record_date)")

        # ECODA activity tracking table
        await db.execute("""
            CREATE TABLE IF NOT EXISTS ecoda_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                ecoda_name TEXT NOT NULL,
                role_type INTEGER DEFAULT 0,
                work_time REAL DEFAULT 0,
                group_id INTEGER,
                team_id INTEGER,
                team_name TEXT,
                record_date TEXT NOT NULL,
                uploaded_by INTEGER,
                upload_timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(ecoda_name, record_date) ON CONFLICT REPLACE
            )
        """)
        # Safe migration check: add team_name column if table already existed without it
        cursor = await db.execute("PRAGMA table_info(ecoda_records)")
        ecoda_cols = [row[1] for row in await cursor.fetchall()]
        if ecoda_cols and "team_name" not in ecoda_cols:
            print("[Schema] Upgrading ecoda_records schema: Adding 'team_name' column...")
            await db.execute("ALTER TABLE ecoda_records ADD COLUMN team_name TEXT")
            print("[Schema] 'team_name' column added successfully!")

        await db.execute("CREATE INDEX IF NOT EXISTS idx_ecoda_name ON ecoda_records(ecoda_name)")
        await db.execute("CREATE INDEX IF NOT EXISTS idx_ecoda_date ON ecoda_records(record_date)")
        await db.execute("CREATE INDEX IF NOT EXISTS idx_ecoda_user ON ecoda_records(user_id)")

        # ECODA excluded workers table (blacklist)
        await db.execute("""
            CREATE TABLE IF NOT EXISTS ecoda_excluded_workers (
                ecoda_name TEXT PRIMARY KEY COLLATE NOCASE,
                added_by INTEGER,
                added_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # ECODA registered teams table
        await db.execute("""
            CREATE TABLE IF NOT EXISTS ecoda_teams (
                name TEXT PRIMARY KEY COLLATE NOCASE,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                is_active INTEGER DEFAULT 1
            )
        """)
        # Seed initial 10 teams if table is empty
        cursor = await db.execute("SELECT COUNT(*) FROM ecoda_teams")
        team_count = (await cursor.fetchone())[0]
        if team_count == 0:
            default_teams = [
                "Delta Force",
                "Nano Banana",
                "Golden Tshushima",
                "Rafael's Cartel [1989]",
                "Night Owls",
                "Totoro",
                "Athena",
                "Flash Point",
                "Rising Horizon",
                "Pixel Hunter",
            ]
            for team_name in default_teams:
                await db.execute(
                    "INSERT OR IGNORE INTO ecoda_teams (name, is_active) VALUES (?, 1)",
                    (team_name,),
                )

        # Settings table for dynamic configurations
        await db.execute("""
            CREATE TABLE IF NOT EXISTS bot_settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        """)

        await db.commit()


async def flush_user_session(user_id: int):
    """Flushes a single user's in-memory session to the database."""
    if user_id not in active_sessions:
        return

    session = active_sessions[user_id]
    if not session.get("channel_id") or not session.get("channel_name"):
        return

    now = time.time()
    elapsed = now - session["last_update"]
    session["last_update"] = now

    if elapsed <= 0:
        return

    state = session["state"]
    unmuted = elapsed if state == "unmuted" else 0
    muted = elapsed if state == "muted" else 0
    deafened = elapsed if state == "deafened" else 0
    record_date = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")

    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute(
            """
            INSERT INTO voice_activity (user_id, user_name, channel_id, channel_name, record_date, unmuted_seconds, muted_seconds, deafened_seconds)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(user_id, channel_id, record_date) DO UPDATE SET
                user_name = excluded.user_name,
                channel_name = excluded.channel_name,
                unmuted_seconds = unmuted_seconds + excluded.unmuted_seconds,
                muted_seconds = muted_seconds + excluded.muted_seconds,
                deafened_seconds = deafened_seconds + excluded.deafened_seconds
            """,
            (user_id, session["name"], session["channel_id"], session["channel_name"], record_date, unmuted, muted, deafened),
        )
        await db.commit()


async def sync_all_sessions():
    """Flushes all active in-memory sessions to the database in a single batch."""
    now = time.time()
    updates = []
    record_date = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")

    for user_id, session in active_sessions.items():
        if not session.get("channel_id") or not session.get("channel_name"):
            continue

        elapsed = now - session["last_update"]
        if elapsed <= 0:
            continue

        session["last_update"] = now
        state = session["state"]
        unmuted = elapsed if state == "unmuted" else 0
        muted = elapsed if state == "muted" else 0
        deafened = elapsed if state == "deafened" else 0

        updates.append((
            user_id, session["name"], session["channel_id"], session["channel_name"],
            record_date, unmuted, muted, deafened
        ))

    if updates:
        async with aiosqlite.connect(DB_FILE) as db:
            await db.executemany(
                """
                INSERT INTO voice_activity (user_id, user_name, channel_id, channel_name, record_date, unmuted_seconds, muted_seconds, deafened_seconds)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(user_id, channel_id, record_date) DO UPDATE SET
                    user_name = excluded.user_name,
                    channel_name = excluded.channel_name,
                    unmuted_seconds = unmuted_seconds + excluded.unmuted_seconds,
                    muted_seconds = muted_seconds + excluded.muted_seconds,
                    deafened_seconds = deafened_seconds + excluded.deafened_seconds
                """,
                updates
            )
            await db.commit()


async def fetch_user_stats(user_id: int, start_date: str, end_date: str) -> dict:
    """Fetches and aggregates per-channel voice stats for a single user."""
    query = """
        SELECT channel_name, SUM(unmuted_seconds), SUM(muted_seconds), SUM(deafened_seconds)
        FROM voice_activity
        WHERE user_id = ? AND record_date BETWEEN ? AND ? AND channel_name IS NOT NULL
        GROUP BY channel_name
    """
    async with aiosqlite.connect(DB_FILE) as db:
        async with db.execute(query, [user_id, start_date, end_date]) as cursor:
            rows = await cursor.fetchall()

    unmuted = sum((r[1] or 0.0) for r in rows) if rows else 0.0
    muted = sum((r[2] or 0.0) for r in rows) if rows else 0.0
    deafened = sum((r[3] or 0.0) for r in rows) if rows else 0.0
    total = unmuted + muted + deafened

    top_vc = "None"
    max_vc_time = 0.0
    channel_breakdown = []

    for r in rows:
        ch_name = r[0]
        if not ch_name:
            continue
        vc_unmuted = r[1] or 0.0
        vc_muted = r[2] or 0.0
        vc_deafened = r[3] or 0.0
        vc_total = vc_unmuted + vc_muted + vc_deafened

        if vc_total > 0:
            channel_breakdown.append((ch_name, vc_total))
        if vc_total > max_vc_time:
            max_vc_time = vc_total
            top_vc = ch_name

    channel_breakdown.sort(key=lambda x: x[1], reverse=True)

    return {
        "unmuted": unmuted,
        "muted": muted,
        "deafened": deafened,
        "total": total,
        "top_vc": top_vc,
        "channel_breakdown": channel_breakdown,
    }


async def fetch_leaderboard_data(start_date: str, end_date: str, limit: int = 500) -> list[dict]:
    """Fetches voice users ranked by total activity for the given date range."""
    query = """
        SELECT user_id, user_name,
               SUM(unmuted_seconds), SUM(muted_seconds), SUM(deafened_seconds),
               SUM(unmuted_seconds + muted_seconds + deafened_seconds) as total
        FROM voice_activity
        WHERE record_date BETWEEN ? AND ?
        GROUP BY user_id
        ORDER BY total DESC
        LIMIT ?
    """
    async with aiosqlite.connect(DB_FILE) as db:
        async with db.execute(query, (start_date, end_date, limit)) as cursor:
            users = await cursor.fetchall()

        leaderboard = []
        for u in users:
            user_id, user_name, unmuted, muted, deafened, total = u
            if total <= 0:
                continue

            ch_query = """
                SELECT channel_name, SUM(unmuted_seconds + muted_seconds + deafened_seconds) as ch_total
                FROM voice_activity
                WHERE user_id = ? AND record_date BETWEEN ? AND ? AND channel_name IS NOT NULL
                GROUP BY channel_name
                ORDER BY ch_total DESC
                LIMIT 1
            """
            async with db.execute(ch_query, (user_id, start_date, end_date)) as ch_cursor:
                ch_row = await ch_cursor.fetchone()
                primary_ch = ch_row[0] if (ch_row and ch_row[0]) else "None"

            leaderboard.append({
                "user_id": user_id,
                "user_name": user_name,
                "unmuted": unmuted or 0.0,
                "muted": muted or 0.0,
                "deafened": deafened or 0.0,
                "total": total or 0.0,
                "primary_channel": primary_ch,
            })

    return leaderboard


async def fetch_role_activity_data(guild: discord.Guild, role: discord.Role, start_date: str, end_date: str) -> dict | None:
    """Fetches voice activity data for all non-bot members of a given role."""
    await sync_all_sessions()

    role_members = [m for m in role.members if not m.bot]
    if not role_members:
        return None

    member_id_map = {m.id: m.display_name for m in role_members}
    member_ids = list(member_id_map.keys())

    user_stats = {
        uid: {
            "user_id": uid,
            "name": member_id_map[uid],
            "unmuted": 0.0,
            "muted": 0.0,
            "deafened": 0.0,
            "total": 0.0,
            "channels": {},
        }
        for uid in member_ids
    }

    chunk_size = 900
    all_rows = []
    async with aiosqlite.connect(DB_FILE) as db:
        for i in range(0, len(member_ids), chunk_size):
            chunk = member_ids[i:i + chunk_size]
            placeholders = ",".join("?" for _ in chunk)
            query = f"""
                SELECT user_id, user_name, channel_name,
                       SUM(unmuted_seconds), SUM(muted_seconds), SUM(deafened_seconds)
                FROM voice_activity
                WHERE user_id IN ({placeholders}) AND record_date BETWEEN ? AND ? AND channel_name IS NOT NULL
                GROUP BY user_id, channel_name
                ORDER BY user_name, channel_name
            """
            params = chunk + [start_date, end_date]
            async with db.execute(query, params) as cursor:
                all_rows.extend(await cursor.fetchall())

    for row in all_rows:
        uid, uname, cname, unmuted_sec, muted_sec, deaf_sec = row
        unmuted_sec = unmuted_sec or 0.0
        muted_sec = muted_sec or 0.0
        deaf_sec = deaf_sec or 0.0
        ch_total = unmuted_sec + muted_sec + deaf_sec

        if uid in user_stats:
            if uname:
                user_stats[uid]["name"] = uname
            user_stats[uid]["unmuted"] += unmuted_sec
            user_stats[uid]["muted"] += muted_sec
            user_stats[uid]["deafened"] += deaf_sec
            user_stats[uid]["total"] += ch_total
            if cname:
                user_stats[uid]["channels"][cname] = {
                    "unmuted": unmuted_sec,
                    "muted": muted_sec,
                    "deafened": deaf_sec,
                    "total": ch_total,
                }

    total_unmuted = sum(u["unmuted"] for u in user_stats.values())
    total_muted = sum(u["muted"] for u in user_stats.values())
    total_deafened = sum(u["deafened"] for u in user_stats.values())
    total_time = total_unmuted + total_muted + total_deafened

    sorted_users = sorted(user_stats.values(), key=lambda x: x["total"], reverse=True)
    active_users = [u for u in sorted_users if u["total"] > 0]
    inactive_users = [u for u in sorted_users if u["total"] == 0]

    return {
        "role": role,
        "total_members": len(role_members),
        "active_users": active_users,
        "inactive_users": inactive_users,
        "total_time": total_time,
        "total_unmuted": total_unmuted,
        "total_muted": total_muted,
        "total_deafened": total_deafened,
        "all_users_sorted": sorted_users,
    }


# ==========================================
# CHART & CSV GENERATORS
# ==========================================
def generate_stats_chart_sync(stats: dict, channel_breakdown: list) -> io.BytesIO:
    """Generates a two-panel matplotlib chart and returns a BytesIO buffer."""
    plt.style.use("dark_background")
    fig = plt.figure(figsize=(12, 6))

    ax1 = fig.add_subplot(121)
    labels = ["Unmuted", "Muted", "Deafened"]
    sizes = [stats["unmuted"], stats["muted"], stats["deafened"]]
    colors = ["#2ecc71", "#f1c40f", "#e74c3c"]

    if sum(sizes) == 0:
        ax1.text(0.5, 0.5, "No Voice Data Available", ha="center", va="center", fontsize=12)
        ax1.axis("off")
    else:
        explode = (0.05, 0.05, 0.05)
        ax1.pie(sizes, explode=explode, labels=labels, colors=colors, autopct="%1.1f%%",
                startangle=90, textprops={"color": "white", "weight": "bold"})
    ax1.set_title("Audio State Distribution", fontsize=14, pad=15)

    ax2 = fig.add_subplot(122)
    valid_channels = [(str(c[0]), c[1]) for c in (channel_breakdown or []) if c and c[0]]
    if not valid_channels:
        ax2.text(0.5, 0.5, "No Channel Data", ha="center", va="center", fontsize=12)
        ax2.axis("off")
    else:
        top_channels = valid_channels[:5]
        c_labels = [c[0][:12] + "..." if len(c[0]) > 12 else c[0] for c in top_channels]
        c_times = [c[1] / 3600 for c in top_channels]

        bars = ax2.bar(c_labels, c_times, color="#3498db", edgecolor="white", linewidth=1)
        ax2.set_ylabel("Hours Spent", fontsize=12)
        ax2.set_title("Top Voice Channels", fontsize=14, pad=15)
        plt.xticks(rotation=30, ha="right")

        for bar in bars:
            yval = bar.get_height()
            ax2.text(bar.get_x() + bar.get_width() / 2, yval + 0.05,
                     f"{yval:.1f}h", ha="center", va="bottom", fontsize=10)

    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format="png", transparent=True, dpi=100)
    plt.close(fig)
    buf.seek(0)
    return buf


async def generate_and_send_csv(target, timeframe_label: str = "last_month"):
    """Generates a CSV voice report and sends it to the given target (channel or interaction)."""
    await sync_all_sessions()
    start_date, end_date = get_date_range(timeframe_label)

    async with aiosqlite.connect(DB_FILE) as db:
        async with db.execute(
            """
            SELECT user_name, user_id, channel_name, SUM(unmuted_seconds), SUM(muted_seconds), SUM(deafened_seconds)
            FROM voice_activity
            WHERE record_date BETWEEN ? AND ? AND channel_name IS NOT NULL
            GROUP BY user_id, channel_name
            ORDER BY user_name, channel_name
            """,
            (start_date, end_date),
        ) as cursor:
            rows = await cursor.fetchall()

    if not rows:
        msg = f"ℹ️ `{start_date}` theke `{end_date}` porjonto kono voice data nei."
        if isinstance(target, discord.Interaction):
            await target.followup.send(msg)
        elif target:
            await target.send(msg)
        return

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "User Name", "User ID", "Voice Channel",
        "Unmuted Time", "Muted Time", "Deafened Time", "Total Time",
        "Unmuted (Sec)", "Muted (Sec)", "Deafened (Sec)", "Total (Sec)"
    ])

    for row in rows:
        uname, uid, cname, unmuted, muted, deaf = row
        total = unmuted + muted + deaf
        writer.writerow([
            uname, uid, cname or "Unknown",
            format_duration(unmuted), format_duration(muted), format_duration(deaf), format_duration(total),
            round(unmuted, 2), round(muted, 2), round(deaf, 2), round(total, 2)
        ])

    output.seek(0)
    file_name = f"Voice_Report_{start_date}_to_{end_date}.csv"
    discord_file = discord.File(fp=io.BytesIO(output.getvalue().encode("utf-8")), filename=file_name)
    content_msg = f"📊 **Detailed Voice Activity Report** ({start_date} to {end_date})"

    if isinstance(target, discord.Interaction):
        await target.followup.send(content=content_msg, file=discord_file)
    elif target:
        await target.send(content=content_msg, file=discord_file)


def generate_role_csv(data: dict, start_date: str, end_date: str) -> discord.File:
    """Generates a CSV file for role activity data."""
    role = data["role"]
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "User Name", "User ID", "Role", "Voice Channel",
        "Unmuted Time", "Muted Time", "Deafened Time", "Total Time",
        "Unmuted (Sec)", "Muted (Sec)", "Deafened (Sec)", "Total (Sec)"
    ])

    for u in data["active_users"]:
        uname = u["name"]
        uid = u["user_id"]
        for cname, ch_data in u["channels"].items():
            writer.writerow([
                uname, uid, role.name, cname,
                format_duration(ch_data["unmuted"]), format_duration(ch_data["muted"]),
                format_duration(ch_data["deafened"]), format_duration(ch_data["total"]),
                round(ch_data["unmuted"], 2), round(ch_data["muted"], 2),
                round(ch_data["deafened"], 2), round(ch_data["total"], 2)
            ])

    for u in data["inactive_users"]:
        writer.writerow([
            u["name"], u["user_id"], role.name, "None (Inactive)",
            "0s", "0s", "0s", "0s",
            0, 0, 0, 0
        ])

    output.seek(0)
    clean_role_name = "".join(c for c in role.name if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    filename = f"Role_{clean_role_name}_{start_date}_to_{end_date}.csv"
    return discord.File(fp=io.BytesIO(output.getvalue().encode("utf-8")), filename=filename)


# ==========================================
# ECODA ACTIVITY TRACKING & HELPERS
# ==========================================
ECODA_CHECKER_ROLE = os.getenv("ECODA_CHECKER_ROLE", "Checker")


async def get_setting(key: str, default: str | None = None) -> str | None:
    """Gets a configuration value from the bot_settings table."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("SELECT value FROM bot_settings WHERE key = ?", (key,))
        row = await cursor.fetchone()
        return row[0] if row else default


async def set_setting(key: str, value: str):
    """Sets or updates a configuration value in the bot_settings table."""
    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute("""
            INSERT INTO bot_settings (key, value) VALUES (?, ?)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """, (key, str(value)))
        await db.commit()


async def is_ecoda_checker(member: discord.Member) -> bool:
    """Returns True if the member has the Checker/Leader role or Administrator permissions."""
    if not isinstance(member, discord.Member):
        return False
    if member.guild_permissions.administrator:
        return True

    # 1. Check dynamically configured role ID from bot_settings
    custom_role_id = await get_setting("ecoda_leader_role_id")
    if custom_role_id and custom_role_id.isdigit():
        target_id = int(custom_role_id)
        if any(r.id == target_id for r in member.roles):
            return True

    # 2. Check role names matching 'checker' or env
    target = ECODA_CHECKER_ROLE.lower()
    for r in member.roles:
        if target in r.name.lower() or "checker" in r.name.lower():
            return True
    return False


def format_hours(hours: float) -> str:
    """Converts decimal hours (e.g. 1.578) into a readable string (e.g. '1h 35m')."""
    if not hours or hours <= 0:
        return "0h 00m"
    total_minutes = int(round(hours * 60))
    h, m = divmod(total_minutes, 60)
    return f"{h}h {m:02d}m"


def find_member_by_name(guild: discord.Guild, name: str) -> discord.Member | None:
    """Finds a member in the guild whose nickname, display_name, or username matches name."""
    if not guild or not name:
        return None
    target = name.strip().lower()

    # 1. Nickname match
    for m in guild.members:
        if m.nick and m.nick.strip().lower() == target:
            return m
    # 2. Display name match
    for m in guild.members:
        if m.display_name.strip().lower() == target:
            return m
    # 3. Global username match
    for m in guild.members:
        if m.name.strip().lower() == target:
            return m
    return None


def resolve_member_display(guild: discord.Guild, ecoda_name: str, cached_user_id: int | None = None) -> tuple[str, bool]:
    """
    Resolves an ECODA worker's display string.
    Returns (display_str, is_matched).
    If matched in Discord server: ('<@user_id>', True)
    If unmatched: ('**ecoda_name**', False)
    """
    if not guild:
        return f"**{ecoda_name}**", False

    if cached_user_id:
        m = guild.get_member(cached_user_id)
        if m:
            return f"<@{m.id}>", True

    m = find_member_by_name(guild, ecoda_name)
    if m:
        return f"<@{m.id}>", True

    return f"**{ecoda_name}**", False


def parse_ecoda_file(file_bytes: bytes, filename: str, default_team_name: str | None = None) -> list[dict]:
    """
    Parses an uploaded CSV or XLSX file and extracts ECODA worker records.
    Returns list of dicts:
      [{"ecoda_name": str, "role_type": int, "work_time": float, "group_id": int|None, "team_id": int|None, "team_name": str|None}]
    """
    ext = filename.lower().split(".")[-1]
    rows_data = []

    if ext in ("xlsx", "xls"):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        iter_rows = list(ws.iter_rows(values_only=True))
        if not iter_rows:
            return []
        headers = [str(h).strip().lower() if h is not None else "" for h in iter_rows[0]]
        for row in iter_rows[1:]:
            if not any(row):
                continue
            row_dict = {}
            for h, v in zip(headers, row):
                if h:
                    row_dict[h] = v
            rows_data.append(row_dict)
    else:
        text = None
        for enc in ("utf-8-sig", "utf-8", "latin1"):
            try:
                text = file_bytes.decode(enc)
                break
            except Exception:
                continue
        if text is None:
            raise ValueError("Could not decode CSV file. Please make sure it is in UTF-8 format.")

        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            clean_row = {k.strip().lower() if k else "": v for k, v in row.items()}
            rows_data.append(clean_row)

    parsed_records = []
    for r in rows_data:
        ecoda_name = (
            r.get("user_name")
            or r.get("username")
            or r.get("worker_name")
            or r.get("name")
            or r.get("ecoda_name")
            or ""
        )
        if not isinstance(ecoda_name, str):
            ecoda_name = str(ecoda_name) if ecoda_name is not None else ""
        ecoda_name = ecoda_name.strip()
        if not ecoda_name:
            continue

        raw_role = r.get("default_role") or r.get("role") or r.get("role_type") or 0
        try:
            role_type = int(raw_role)
        except (ValueError, TypeError):
            role_type = 0

        raw_time = r.get("work_time") or r.get("work_hour") or r.get("hours") or r.get("time") or 0
        try:
            work_time = float(raw_time)
        except (ValueError, TypeError):
            work_time = 0.0

        raw_group = r.get("group_id") or r.get("group")
        group_id = int(raw_group) if raw_group is not None and str(raw_group).isdigit() else None

        raw_team = r.get("team_id") or r.get("team")
        team_id = int(raw_team) if raw_team is not None and str(raw_team).isdigit() else None

        # Team name extraction (from column or default parameter)
        raw_team_name = r.get("team_name") or r.get("team_title")
        if not raw_team_name and "team" in r and not str(r["team"]).isdigit():
            raw_team_name = r["team"]

        team_name_str = None
        if raw_team_name is not None and str(raw_team_name).strip():
            team_name_str = str(raw_team_name).strip()
        elif default_team_name and default_team_name.strip():
            team_name_str = default_team_name.strip()

        parsed_records.append({
            "ecoda_name": ecoda_name,
            "role_type": role_type,
            "work_time": work_time,
            "group_id": group_id,
            "team_id": team_id,
            "team_name": team_name_str,
        })

    return parsed_records


async def save_ecoda_records(records: list[dict], record_date: str, uploaded_by: int) -> int:
    """Inserts or updates ECODA worker records for the given date, skipping blacklisted workers."""
    if not records:
        return 0
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("SELECT LOWER(TRIM(ecoda_name)) FROM ecoda_excluded_workers")
        excluded_set = {row[0] for row in await cursor.fetchall()}

        saved_count = 0
        for r in records:
            if r["ecoda_name"].strip().lower() in excluded_set:
                continue
            await db.execute("""
                INSERT INTO ecoda_records (
                    user_id, ecoda_name, role_type, work_time, group_id, team_id, team_name, record_date, uploaded_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(ecoda_name, record_date) DO UPDATE SET
                    user_id = excluded.user_id,
                    role_type = excluded.role_type,
                    work_time = excluded.work_time,
                    group_id = excluded.group_id,
                    team_id = excluded.team_id,
                    team_name = COALESCE(excluded.team_name, ecoda_records.team_name),
                    uploaded_by = excluded.uploaded_by,
                    upload_timestamp = CURRENT_TIMESTAMP
            """, (
                r.get("user_id"),
                r["ecoda_name"],
                r.get("role_type", 0),
                r.get("work_time", 0.0),
                r.get("group_id"),
                r.get("team_id"),
                r.get("team_name"),
                record_date,
                uploaded_by,
            ))
            saved_count += 1
        await db.commit()
    return saved_count


async def fetch_ecoda_leaderboard_data(
    start_date: str | None = None,
    end_date: str | None = None,
    role_filter: int | None = None,
    status_filter: str = "active",
    limit: int = 500,
) -> list[dict]:
    """
    Fetches aggregated ECODA work hour records.
    status_filter: 'active' (hours > 0), 'inactive' (hours == 0), or 'all'
    """
    query = """
        SELECT 
            ecoda_name,
            MAX(user_id) as user_id,
            role_type,
            MAX(team_name) as team_name,
            SUM(work_time) as total_hours,
            COUNT(DISTINCT record_date) as active_days
        FROM ecoda_records
        WHERE 1=1
          AND LOWER(TRIM(ecoda_name)) NOT IN (SELECT LOWER(TRIM(ecoda_name)) FROM ecoda_excluded_workers)
    """
    params = []
    if start_date:
        query += " AND record_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND record_date <= ?"
        params.append(end_date)
    if role_filter is not None:
        query += " AND role_type = ?"
        params.append(role_filter)

    query += " GROUP BY ecoda_name"

    if status_filter == "active":
        query += " HAVING total_hours > 0.0001 ORDER BY total_hours DESC"
    elif status_filter == "inactive":
        query += " HAVING total_hours <= 0.0001 ORDER BY team_name ASC, ecoda_name ASC"
    else:
        query += " ORDER BY total_hours DESC, ecoda_name ASC"

    query += " LIMIT ?"
    params.append(limit)

    async with aiosqlite.connect(DB_FILE) as db:
        db.row_factory = aiosqlite.Row
        cursor = await db.execute(query, params)
        rows = await cursor.fetchall()
        return [dict(r) for r in rows]


async def update_ecoda_work_time(
    worker_query: str,
    record_date: str,
    new_hours: float,
    updater_id: int,
    team_name: str | None = None,
) -> dict | None:
    """
    Finds a record by worker name (or user_id) on record_date and updates work_time (and optionally team_name).
    Returns a dict with updated record details, or None if not found.
    """
    clean_query = worker_query.strip()
    digits = re.findall(r"\d+", clean_query)
    user_id = int(digits[0]) if digits and len(digits[0]) >= 15 else None

    async with aiosqlite.connect(DB_FILE) as db:
        db.row_factory = aiosqlite.Row
        row = None

        if user_id:
            cursor = await db.execute(
                "SELECT * FROM ecoda_records WHERE user_id = ? AND record_date = ?",
                (user_id, record_date),
            )
            row = await cursor.fetchone()

        if not row:
            cursor = await db.execute(
                "SELECT * FROM ecoda_records WHERE LOWER(ecoda_name) = ? AND record_date = ?",
                (clean_query.lower(), record_date),
            )
            row = await cursor.fetchone()

        if not row:
            cursor = await db.execute(
                "SELECT * FROM ecoda_records WHERE LOWER(ecoda_name) LIKE ? AND record_date = ? LIMIT 1",
                (f"%{clean_query.lower()}%", record_date),
            )
            row = await cursor.fetchone()

        if not row:
            return None

        record_id = row["id"]
        old_hours = row["work_time"]
        ecoda_name = row["ecoda_name"]
        matched_uid = row["user_id"]
        existing_team = row["team_name"]
        final_team = team_name.strip() if team_name and team_name.strip() else existing_team

        if team_name is not None and team_name.strip():
            await db.execute("""
                UPDATE ecoda_records
                SET work_time = ?, team_name = ?, uploaded_by = ?, upload_timestamp = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (new_hours, final_team, updater_id, record_id))
        else:
            await db.execute("""
                UPDATE ecoda_records
                SET work_time = ?, uploaded_by = ?, upload_timestamp = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (new_hours, updater_id, record_id))
        await db.commit()

        return {
            "id": record_id,
            "ecoda_name": ecoda_name,
            "user_id": matched_uid,
            "old_hours": old_hours,
            "new_hours": new_hours,
            "team_name": final_team,
            "record_date": record_date,
        }


async def add_ecoda_excluded_worker(ecoda_name: str, added_by: int) -> bool:
    """Adds a worker to the excluded/blacklist table to hide them from all leaderboards."""
    clean_name = ecoda_name.strip()
    if not clean_name:
        return False
    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute("""
            INSERT INTO ecoda_excluded_workers (ecoda_name, added_by)
            VALUES (?, ?)
            ON CONFLICT(ecoda_name) DO UPDATE SET added_by = excluded.added_by, added_at = CURRENT_TIMESTAMP
        """, (clean_name, added_by))
        await db.commit()
    return True


async def remove_ecoda_excluded_worker(ecoda_name: str) -> bool:
    """Removes a worker from the excluded/blacklist table."""
    clean_name = ecoda_name.strip()
    if not clean_name:
        return False
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute(
            "DELETE FROM ecoda_excluded_workers WHERE LOWER(TRIM(ecoda_name)) = LOWER(TRIM(?))",
            (clean_name,)
        )
        await db.commit()
        return cursor.rowcount > 0


async def get_ecoda_excluded_workers() -> list[dict]:
    """Returns a list of all blacklisted workers."""
    async with aiosqlite.connect(DB_FILE) as db:
        db.row_factory = aiosqlite.Row
        cursor = await db.execute(
            "SELECT ecoda_name, added_by, added_at FROM ecoda_excluded_workers ORDER BY ecoda_name ASC"
        )
        rows = await cursor.fetchall()
        return [dict(r) for r in rows]


# ==========================================
# ECODA TEAM MANAGEMENT
# ==========================================
async def get_all_teams() -> list[str]:
    """Returns a list of all active registered team names sorted alphabetically."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute(
            "SELECT name FROM ecoda_teams WHERE is_active = 1 ORDER BY name COLLATE NOCASE ASC"
        )
        rows = await cursor.fetchall()
        return [row[0] for row in rows]


async def add_team(name: str) -> tuple[bool, str]:
    """Adds a new team to ecoda_teams. Returns (success, message)."""
    clean_name = name.strip()
    if not clean_name:
        return False, "❌ Team name cannot be empty."
    if len(clean_name) > 50:
        return False, "❌ Team name must be 50 characters or less."

    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute(
            "SELECT is_active FROM ecoda_teams WHERE LOWER(name) = LOWER(?)",
            (clean_name,),
        )
        row = await cursor.fetchone()
        if row:
            if row[0] == 1:
                return False, f"⚠️ Team **`{clean_name}`** already exists in the team roster!"
            else:
                await db.execute(
                    "UPDATE ecoda_teams SET is_active = 1 WHERE LOWER(name) = LOWER(?)",
                    (clean_name,),
                )
                await db.commit()
                return True, f"✅ Team **`{clean_name}`** has been reactivated successfully!"

        await db.execute(
            "INSERT INTO ecoda_teams (name, is_active) VALUES (?, 1)",
            (clean_name,),
        )
        await db.commit()
        return True, f"✅ Team **`{clean_name}`** has been added to the team roster!"


async def remove_team(name: str) -> tuple[bool, str]:
    """Marks a team as inactive. Returns (success, message)."""
    clean_name = name.strip()
    if not clean_name:
        return False, "❌ Team name cannot be empty."

    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute(
            "SELECT name FROM ecoda_teams WHERE LOWER(name) = LOWER(?) AND is_active = 1",
            (clean_name,),
        )
        row = await cursor.fetchone()
        if not row:
            return False, f"⚠️ Team **`{clean_name}`** was not found in active teams."

        actual_name = row[0]
        await db.execute(
            "UPDATE ecoda_teams SET is_active = 0 WHERE LOWER(name) = LOWER(?)",
            (clean_name,),
        )
        await db.commit()
        return True, f"✅ Team **`{actual_name}`** has been removed from active teams."


async def get_team_details_list() -> list[dict]:
    """Returns all registered teams with record stats."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("""
            SELECT 
                t.name,
                t.is_active,
                t.created_at,
                COUNT(DISTINCT r.ecoda_name) as worker_count,
                MAX(r.record_date) as last_record_date
            FROM ecoda_teams t
            LEFT JOIN ecoda_records r ON LOWER(TRIM(t.name)) = LOWER(TRIM(r.team_name))
            GROUP BY t.name
            ORDER BY t.is_active DESC, t.name COLLATE NOCASE ASC
        """)
        rows = await cursor.fetchall()
        return [
            {
                "name": row[0],
                "is_active": bool(row[1]),
                "created_at": row[2],
                "worker_count": row[3] or 0,
                "last_record_date": row[4],
            }
            for row in rows
        ]


async def fetch_team_upload_status(record_date: str) -> dict:
    """
    Fetches upload status for all registered teams for a specific date.
    Returns summary stats and breakdown of uploaded vs pending teams.
    """
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute(
            "SELECT name FROM ecoda_teams WHERE is_active = 1 ORDER BY name COLLATE NOCASE ASC"
        )
        all_teams = [row[0] for row in await cursor.fetchall()]

        cursor = await db.execute("""
            SELECT 
                TRIM(team_name) as t_name,
                COUNT(*) as worker_count,
                SUM(work_time) as total_hours,
                MAX(upload_timestamp) as last_upload,
                MAX(uploaded_by) as uploaded_by
            FROM ecoda_records
            WHERE record_date = ? 
              AND team_name IS NOT NULL 
              AND TRIM(team_name) != ''
            GROUP BY LOWER(TRIM(team_name))
        """, (record_date,))
        rows = await cursor.fetchall()

    upload_map = {}
    for row in rows:
        t_name = row[0]
        upload_map[t_name.lower()] = {
            "team_name": t_name,
            "worker_count": row[1],
            "total_hours": row[2] or 0.0,
            "last_upload": row[3],
            "uploaded_by": row[4],
        }

    uploaded = []
    pending = []

    for t in all_teams:
        t_lower = t.lower()
        if t_lower in upload_map:
            info = upload_map.pop(t_lower)
            info["team_name"] = t
            uploaded.append(info)
        else:
            pending.append(t)

    for remaining in upload_map.values():
        uploaded.append(remaining)

    total_teams = len(uploaded) + len(pending)
    uploaded_count = len(uploaded)
    pending_count = len(pending)
    percentage = int(round((uploaded_count / total_teams) * 100)) if total_teams > 0 else 0

    return {
        "record_date": record_date,
        "total_teams": total_teams,
        "uploaded_count": uploaded_count,
        "pending_count": pending_count,
        "percentage": percentage,
        "uploaded": uploaded,
        "pending": pending,
    }


async def add_ecoda_manual_record(
    worker_name: str,
    hours: float,
    record_date: str,
    role_type: int = 0,
    team_name: str | None = None,
    user_id: int | None = None,
    added_by: int = 0,
) -> dict:
    """Manually creates or updates an ECODA work record for a worker on a specific date."""
    clean_name = worker_name.strip()
    clean_team = team_name.strip() if team_name and team_name.strip() else None
    hours = max(0.0, float(hours))

    async with aiosqlite.connect(DB_FILE) as db:
        await db.execute("""
            INSERT INTO ecoda_records (
                user_id, ecoda_name, role_type, work_time, team_name, record_date, uploaded_by
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(ecoda_name, record_date) DO UPDATE SET
                user_id = COALESCE(excluded.user_id, ecoda_records.user_id),
                role_type = excluded.role_type,
                work_time = excluded.work_time,
                team_name = COALESCE(excluded.team_name, ecoda_records.team_name),
                uploaded_by = excluded.uploaded_by,
                upload_timestamp = CURRENT_TIMESTAMP
        """, (user_id, clean_name, role_type, hours, clean_team, record_date, added_by))
        await db.commit()

    return {
        "ecoda_name": clean_name,
        "user_id": user_id,
        "role_type": role_type,
        "hours": hours,
        "team_name": clean_team,
        "record_date": record_date,
    }


async def delete_ecoda_record(
    worker_query: str,
    record_date: str | None = None,
) -> tuple[int, str]:
    """
    Deletes ECODA records matching worker_query (user_id or name).
    If record_date is provided and not in ('all', None), deletes only for that date.
    Returns (deleted_count, matched_display_name).
    """
    clean_query = worker_query.strip()
    digits = re.findall(r"\d+", clean_query)
    user_id = int(digits[0]) if digits and len(digits[0]) >= 15 else None

    async with aiosqlite.connect(DB_FILE) as db:
        db.row_factory = aiosqlite.Row

        # Try to find existing record to get canonical name
        row = None
        if user_id:
            cursor = await db.execute("SELECT ecoda_name FROM ecoda_records WHERE user_id = ? LIMIT 1", (user_id,))
            row = await cursor.fetchone()
        if not row:
            cursor = await db.execute("SELECT ecoda_name FROM ecoda_records WHERE LOWER(TRIM(ecoda_name)) = LOWER(TRIM(?)) LIMIT 1", (clean_query,))
            row = await cursor.fetchone()
        if not row:
            cursor = await db.execute("SELECT ecoda_name FROM ecoda_records WHERE LOWER(ecoda_name) LIKE ? LIMIT 1", (f"%{clean_query.lower()}%",))
            row = await cursor.fetchone()

        matched_name = row["ecoda_name"] if row else clean_query

        conditions = []
        params = []

        if user_id:
            conditions.append("(user_id = ? OR LOWER(TRIM(ecoda_name)) = LOWER(TRIM(?)))")
            params.extend([user_id, clean_query])
        else:
            if row:
                conditions.append("LOWER(TRIM(ecoda_name)) = LOWER(TRIM(?))")
                params.append(matched_name)
            else:
                conditions.append("LOWER(ecoda_name) LIKE ?")
                params.append(f"%{clean_query.lower()}%")

        if record_date and record_date.lower() != "all":
            conditions.append("record_date = ?")
            params.append(record_date)

        where_clause = " AND ".join(conditions)
        cursor = await db.execute(f"DELETE FROM ecoda_records WHERE {where_clause}", params)
        await db.commit()
        return cursor.rowcount, matched_name


async def delete_ecoda_by_date(record_date: str) -> int:
    """Deletes all ECODA records on a specific date. Returns number of rows deleted."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("DELETE FROM ecoda_records WHERE record_date = ?", (record_date,))
        await db.commit()
        return cursor.rowcount


async def reset_ecoda_records() -> int:
    """Deletes all records from ecoda_records table. Preserves voice_activity and settings."""
    async with aiosqlite.connect(DB_FILE) as db:
        cursor = await db.execute("DELETE FROM ecoda_records")
        await db.commit()
        return cursor.rowcount



# ==========================================
# LIVE LEADERBOARD IN-MEMORY STATE & VIEWS
# ==========================================
live_board_state = {
    "voice_page": 1,
    "voice_cutoff": "current",
    "ecoda_page": 1,
    "ecoda_cutoff": "current",
    "ecoda_role": 0,           # 0: Labelers, 1: Checkers, None: All Roles
    "ecoda_status": "active",  # 'active' (work_time > 0) or 'inactive' (work_time == 0)
    "status_date": datetime.now(LOCAL_TZ).strftime("%Y-%m-%d"),
}


def build_live_voice_leaderboard_embed(
    guild: discord.Guild,
    data: list[dict],
    start_date: str,
    end_date: str,
    cutoff_label: str,
    page: int = 1,
    page_size: int = 10,
) -> tuple[discord.Embed, int, int]:
    """Builds the paginated embed for the persistent live Voice leaderboard."""
    total_items = len(data)
    total_pages = max(1, (total_items + page_size - 1) // page_size)
    page = min(max(1, page), total_pages)

    embed = discord.Embed(
        title="🏆 Live Discord Voice Activity Leaderboard",
        description=(
            f"🔴 **LIVE** • Auto-updates every 5 minutes\n"
            f"📅 **Timeframe:** `{cutoff_label}` ({start_date} to {end_date})\n\n"
        ),
        color=discord.Color.gold(),
    )

    if not data:
        embed.description += "ℹ️ *No voice activity recorded for this cutoff yet.*"
    else:
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        page_entries = data[start_idx:end_idx]

        medals = ["🥇", "🥈", "🥉"]
        for idx, entry in enumerate(page_entries, start=start_idx + 1):
            rank_str = medals[idx - 1] if idx <= 3 else f"`#{idx}`"
            name = entry["user_name"]
            total_str = format_duration(entry["total"])
            unmuted_str = format_duration(entry["unmuted"])
            muted_str = format_duration(entry["muted"])
            line = (
                f"{rank_str} **{name}** — `{total_str}`\n"
                f"└ 🟢 `{unmuted_str}` | 🟡 `{muted_str}`\n"
            )
            embed.description += line

    embed.set_footer(text=f"Page {page} of {total_pages} • Total: {total_items} Active Voice Members • Voice Tracker")
    embed.timestamp = datetime.now(LOCAL_TZ)
    return embed, page, total_pages


def build_live_ecoda_leaderboard_embed(
    guild: discord.Guild,
    data: list[dict],
    start_date: str,
    end_date: str,
    cutoff_label: str,
    role_filter: int | None = 0,
    status_filter: str = "active",
    page: int = 1,
    page_size: int = 10,
) -> tuple[discord.Embed, int, int]:
    """Builds the paginated embed for the persistent live ECODA leaderboard with separate role & active/inactive views."""
    total_items = len(data)
    total_pages = max(1, (total_items + page_size - 1) // page_size)
    page = min(max(1, page), total_pages)

    if role_filter == 0:
        role_label = "🏷️ Labelers"
    elif role_filter == 1:
        role_label = "🔍 Checkers"
    else:
        role_label = "👥 All Roles"

    if status_filter == "inactive":
        embed = discord.Embed(
            title=f"💼 Live ECODA Leaderboard — 🔴 Inactive Members ({role_label})",
            description=(
                f"🔴 **INACTIVE MEMBERS (0 Hours Worked)** • Auto-updates on upload\n"
                f"📅 **Timeframe:** `{cutoff_label}` ({start_date} to {end_date})\n"
                f"🎭 **Category:** `{role_label}`\n"
                f"ℹ️ *Eder ei cutoff period e kono work hour record hoyni (0h).*\n\n"
            ),
            color=discord.Color.red(),
        )
    else:
        embed = discord.Embed(
            title=f"💼 Live ECODA Leaderboard — {role_label}",
            description=(
                f"🔴 **LIVE** • Auto-updates on upload & sync\n"
                f"📅 **Timeframe:** `{cutoff_label}` ({start_date} to {end_date})\n"
                f"🎭 **Ranking:** `{role_label}` (Separated Rankings)\n\n"
            ),
            color=discord.Color.teal() if role_filter == 0 else discord.Color.blue(),
        )

    if not data:
        if status_filter == "inactive":
            embed.description += f"🎉 *Shobai active! Ei cutoff period e kono inactive {role_label} nei.*"
        else:
            embed.description += f"ℹ️ *No {role_label} work hours recorded for this cutoff yet.*"
    else:
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        page_entries = data[start_idx:end_idx]

        medals = ["🥇", "🥈", "🥉"]
        for idx, item in enumerate(page_entries, start=start_idx + 1):
            ecoda_name = item["ecoda_name"]
            user_id = item["user_id"]
            role_type = item["role_type"]
            total_hours = item["total_hours"]
            active_days = item["active_days"]
            team_name = item.get("team_name")

            display_name, _ = resolve_member_display(guild, ecoda_name, user_id)
            team_badge = f" `[{team_name}]`" if team_name else ""

            if status_filter == "inactive":
                line = f"`#{idx}` {display_name}{team_badge} — **`0h 00m`** ⚠️ *Inactive*\n"
            else:
                rank_str = medals[idx - 1] if idx <= 3 else f"`#{idx}`"
                time_formatted = format_hours(total_hours)
                role_icon = "🏷️" if role_type == 0 else "🔍"
                days_str = f"{active_days}d"
                line = f"{rank_str} {display_name}{team_badge} — **`{time_formatted}`** (`{role_icon}` • `{days_str}`)\n"

            embed.description += line

    status_tag = f"Inactive Workers (0h)" if status_filter == "inactive" else f"Active {role_label}"
    embed.set_footer(text=f"Page {page} of {total_pages} • Total: {total_items} {status_tag} • ECODA Tracker")
    embed.timestamp = datetime.now(LOCAL_TZ)
    return embed, page, total_pages


class LiveVoiceLeaderboardView(discord.ui.View):
    def __init__(self, page: int = 1, total_pages: int = 1, cutoff_mode: str = "current"):
        super().__init__(timeout=None)
        self.update_buttons(page, total_pages, cutoff_mode)

    def update_buttons(self, page: int, total_pages: int, cutoff_mode: str):
        self.btn_current.style = discord.ButtonStyle.primary if cutoff_mode == "current" else discord.ButtonStyle.secondary
        self.btn_prev_cutoff.style = discord.ButtonStyle.primary if cutoff_mode == "previous" else discord.ButtonStyle.secondary
        self.btn_prev_page.disabled = (page <= 1)
        self.btn_indicator.label = f"📄 {page} / {max(1, total_pages)}"
        self.btn_next_page.disabled = (page >= total_pages)

    @discord.ui.button(label="⏳ Current Cutoff", style=discord.ButtonStyle.primary, custom_id="live_voice:cutoff_curr", row=0)
    async def btn_current(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["voice_cutoff"] = "current"
        live_board_state["voice_page"] = 1
        await interaction.response.defer()
        await update_live_voice_leaderboard(interaction.client)

    @discord.ui.button(label="⏪ Previous Cutoff", style=discord.ButtonStyle.secondary, custom_id="live_voice:cutoff_prev", row=0)
    async def btn_prev_cutoff(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["voice_cutoff"] = "previous"
        live_board_state["voice_page"] = 1
        await interaction.response.defer()
        await update_live_voice_leaderboard(interaction.client)

    @discord.ui.button(label="◀️ Prev", style=discord.ButtonStyle.secondary, custom_id="live_voice:prev", row=1)
    async def btn_prev_page(self, interaction: discord.Interaction, button: discord.ui.Button):
        if live_board_state["voice_page"] > 1:
            live_board_state["voice_page"] -= 1
        await interaction.response.defer()
        await update_live_voice_leaderboard(interaction.client)

    @discord.ui.button(label="📄 1 / 1", style=discord.ButtonStyle.secondary, disabled=True, custom_id="live_voice:indicator", row=1)
    async def btn_indicator(self, interaction: discord.Interaction, button: discord.ui.Button):
        pass

    @discord.ui.button(label="Next ▶️", style=discord.ButtonStyle.secondary, custom_id="live_voice:next", row=1)
    async def btn_next_page(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["voice_page"] += 1
        await interaction.response.defer()
        await update_live_voice_leaderboard(interaction.client)

    @discord.ui.button(label="🔄", style=discord.ButtonStyle.primary, custom_id="live_voice:refresh", row=1)
    async def btn_refresh(self, interaction: discord.Interaction, button: discord.ui.Button):
        await sync_all_sessions()
        await interaction.response.defer()
        await update_live_voice_leaderboard(interaction.client)


class LiveEcodaLeaderboardView(discord.ui.View):
    def __init__(
        self,
        page: int = 1,
        total_pages: int = 1,
        cutoff_mode: str = "current",
        role_filter: int | None = 0,
        status_filter: str = "active",
    ):
        super().__init__(timeout=None)
        self.update_buttons(page, total_pages, cutoff_mode, role_filter, status_filter)

    def update_buttons(
        self,
        page: int,
        total_pages: int,
        cutoff_mode: str,
        role_filter: int | None,
        status_filter: str,
    ):
        # Row 0: Role Selection (Separated Rankings)
        self.btn_labelers.style = (
            discord.ButtonStyle.primary if role_filter == 0 else discord.ButtonStyle.secondary
        )
        self.btn_checkers.style = (
            discord.ButtonStyle.primary if role_filter == 1 else discord.ButtonStyle.secondary
        )
        self.btn_all_roles.style = (
            discord.ButtonStyle.primary if role_filter is None else discord.ButtonStyle.secondary
        )

        # Row 1: Status Selection (Active vs Inactive) + Refresh
        self.btn_active.style = (
            discord.ButtonStyle.success if status_filter == "active" else discord.ButtonStyle.secondary
        )
        self.btn_inactive.style = (
            discord.ButtonStyle.danger if status_filter == "inactive" else discord.ButtonStyle.secondary
        )

        # Row 2: Cutoff Selection
        self.btn_current.style = (
            discord.ButtonStyle.primary if cutoff_mode == "current" else discord.ButtonStyle.secondary
        )
        self.btn_prev_cutoff.style = (
            discord.ButtonStyle.primary if cutoff_mode == "previous" else discord.ButtonStyle.secondary
        )

        # Row 3: Pagination
        self.btn_prev_page.disabled = (page <= 1)
        self.btn_indicator.label = f"📄 {page} / {max(1, total_pages)}"
        self.btn_next_page.disabled = (page >= total_pages)

    @discord.ui.button(label="🏷️ Labelers", style=discord.ButtonStyle.primary, custom_id="live_ecoda:role_labeler", row=0)
    async def btn_labelers(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_role"] = 0
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="🔍 Checkers", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:role_checker", row=0)
    async def btn_checkers(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_role"] = 1
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="👥 All Roles", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:role_all", row=0)
    async def btn_all_roles(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_role"] = None
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="🟢 Active Members", style=discord.ButtonStyle.success, custom_id="live_ecoda:status_active", row=1)
    async def btn_active(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_status"] = "active"
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="🔴 Inactive (0h)", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:status_inactive", row=1)
    async def btn_inactive(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_status"] = "inactive"
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="🔄 Refresh", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:refresh", row=1)
    async def btn_refresh(self, interaction: discord.Interaction, button: discord.ui.Button):
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="⏳ Current Cutoff", style=discord.ButtonStyle.primary, custom_id="live_ecoda:cutoff_curr", row=2)
    async def btn_current(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_cutoff"] = "current"
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="⏪ Previous Cutoff", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:cutoff_prev", row=2)
    async def btn_prev_cutoff(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_cutoff"] = "previous"
        live_board_state["ecoda_page"] = 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="◀️ Prev", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:prev", row=3)
    async def btn_prev_page(self, interaction: discord.Interaction, button: discord.ui.Button):
        if live_board_state.get("ecoda_page", 1) > 1:
            live_board_state["ecoda_page"] -= 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)

    @discord.ui.button(label="📄 1 / 1", style=discord.ButtonStyle.secondary, disabled=True, custom_id="live_ecoda:indicator", row=3)
    async def btn_indicator(self, interaction: discord.Interaction, button: discord.ui.Button):
        pass

    @discord.ui.button(label="Next ▶️", style=discord.ButtonStyle.secondary, custom_id="live_ecoda:next", row=3)
    async def btn_next_page(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["ecoda_page"] = live_board_state.get("ecoda_page", 1) + 1
        await interaction.response.defer()
        await update_live_ecoda_leaderboard(interaction.client)


async def update_live_voice_leaderboard(bot: commands.Bot):
    """Updates the live Voice leaderboard message in the designated channel."""
    channel_id_str = await get_setting("live_leaderboard_channel_id")
    voice_msg_id_str = await get_setting("live_voice_msg_id")
    if not channel_id_str or not channel_id_str.isdigit() or not voice_msg_id_str:
        return

    channel = bot.get_channel(int(channel_id_str))
    if not channel:
        try:
            channel = await bot.fetch_channel(int(channel_id_str))
        except Exception:
            return

    cutoff_mode = live_board_state.get("voice_cutoff", "current")
    start_date, end_date, cutoff_label = get_cutoff_dates(cutoff_mode)
    voice_data = await fetch_leaderboard_data(start_date, end_date, limit=500)

    req_page = live_board_state.get("voice_page", 1)
    voice_embed, act_page, total_pages = build_live_voice_leaderboard_embed(
        channel.guild, voice_data, start_date, end_date, cutoff_label, page=req_page, page_size=10
    )
    live_board_state["voice_page"] = act_page

    view = LiveVoiceLeaderboardView(page=act_page, total_pages=total_pages, cutoff_mode=cutoff_mode)
    try:
        voice_msg = await channel.fetch_message(int(voice_msg_id_str))
        await voice_msg.edit(content="", embed=voice_embed, view=view)
    except Exception as e:
        print(f"Notice: Failed to update live voice message: {e}")


async def update_live_ecoda_leaderboard(bot: commands.Bot):
    """Updates the live ECODA leaderboard message in the designated channel."""
    channel_id_str = await get_setting("live_leaderboard_channel_id")
    ecoda_msg_id_str = await get_setting("live_ecoda_msg_id")
    if not channel_id_str or not channel_id_str.isdigit() or not ecoda_msg_id_str:
        return

    channel = bot.get_channel(int(channel_id_str))
    if not channel:
        try:
            channel = await bot.fetch_channel(int(channel_id_str))
        except Exception:
            return

    cutoff_mode = live_board_state.get("ecoda_cutoff", "current")
    role_filter = live_board_state.get("ecoda_role", 0)  # default: Labelers (0)
    status_filter = live_board_state.get("ecoda_status", "active")  # default: active
    start_date, end_date, cutoff_label = get_cutoff_dates(cutoff_mode)

    ecoda_data = await fetch_ecoda_leaderboard_data(
        start_date=start_date,
        end_date=end_date,
        role_filter=role_filter,
        status_filter=status_filter,
        limit=500,
    )

    req_page = live_board_state.get("ecoda_page", 1)
    ecoda_embed, act_page, total_pages = build_live_ecoda_leaderboard_embed(
        channel.guild,
        ecoda_data,
        start_date,
        end_date,
        cutoff_label,
        role_filter=role_filter,
        status_filter=status_filter,
        page=req_page,
        page_size=10,
    )
    live_board_state["ecoda_page"] = act_page

    view = LiveEcodaLeaderboardView(
        page=act_page,
        total_pages=total_pages,
        cutoff_mode=cutoff_mode,
        role_filter=role_filter,
        status_filter=status_filter,
    )
    try:
        ecoda_msg = await channel.fetch_message(int(ecoda_msg_id_str))
        await ecoda_msg.edit(content="", embed=ecoda_embed, view=view)
    except Exception as e:
        print(f"Notice: Failed to update live ecoda message: {e}")


def build_team_upload_status_embed(
    guild: discord.Guild | None,
    status_data: dict,
    record_date: str,
) -> discord.Embed:
    """Builds the rich embed for ECODA Daily Team Upload Status."""
    total_teams = status_data["total_teams"]
    uploaded_count = status_data["uploaded_count"]
    pending_count = status_data["pending_count"]
    percentage = status_data["percentage"]
    uploaded = status_data["uploaded"]
    pending = status_data["pending"]

    today_str = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
    try:
        dt_obj = datetime.strptime(record_date, "%Y-%m-%d")
        formatted_date = dt_obj.strftime("%A, %B %d, %Y")
    except Exception:
        formatted_date = record_date

    if record_date == today_str:
        date_badge = "✨ **Today**"
    elif record_date < today_str:
        date_badge = "⏳ **Past Date**"
    else:
        date_badge = "🔮 **Future Date**"

    filled = int(round((percentage / 100) * 10))
    bar = "🟩" * filled + "⬜" * (10 - filled)

    if pending_count == 0 and total_teams > 0:
        color = discord.Color.green()
    elif uploaded_count > 0:
        color = discord.Color.gold()
    else:
        color = discord.Color.red()

    embed = discord.Embed(
        title="📊 ECODA Daily Team Upload Status",
        color=color,
    )

    desc_lines = [
        f"📅 **Date:** `{record_date}` ({date_badge}) — *{formatted_date}*",
        f"**Progress:** {bar} **{percentage}%**",
        f"**Overview:** ✅ **`{uploaded_count}`** Uploaded • ❌ **`{pending_count}`** Pending • Total: **`{total_teams}`** Teams\n",
    ]
    embed.description = "\n".join(desc_lines)

    if uploaded:
        up_lines = []
        for u in uploaded:
            t_name = u["team_name"]
            w_count = u["worker_count"]
            h_str = format_hours(u["total_hours"])
            u_by = f"<@{u['uploaded_by']}>" if u.get("uploaded_by") else "*Sheet Upload*"

            raw_ts = u.get("last_upload")
            time_part = ""
            if raw_ts:
                try:
                    ts_dt = datetime.fromisoformat(str(raw_ts).replace("Z", ""))
                    time_part = f" • 🕒 `{ts_dt.strftime('%I:%M %p')}`"
                except Exception:
                    pass

            up_lines.append(f"🟢 **{t_name}**\n  └ {u_by} • ⏱️ `{h_str}` ({w_count} workers){time_part}")

        chunk = "\n".join(up_lines)
        if len(chunk) > 1024:
            chunk = chunk[:1020] + "..."
        embed.add_field(name=f"✅ Uploaded Teams ({uploaded_count})", value=chunk, inline=False)
    else:
        embed.add_field(
            name=f"✅ Uploaded Teams ({uploaded_count})",
            value="*No team sheets uploaded yet for this date.*",
            inline=False,
        )

    if pending:
        pend_lines = [f"🔴 **{t}** — ⚠️ *Awaiting sheet upload*" for t in pending]
        chunk = "\n".join(pend_lines)
        if len(chunk) > 1024:
            chunk = chunk[:1020] + "..."
        embed.add_field(name=f"❌ Pending Teams ({pending_count})", value=chunk, inline=False)
    else:
        embed.add_field(
            name=f"❌ Pending Teams (0)",
            value="🎉 *All registered teams have uploaded their sheets for this date!*",
            inline=False,
        )

    embed.set_footer(text="ECODA Operational Tracker • Auto-refreshes on upload")
    embed.timestamp = datetime.now(LOCAL_TZ)
    return embed


class LiveTeamUploadStatusView(discord.ui.View):
    def __init__(self, record_date: str | None = None):
        super().__init__(timeout=None)
        date_to_use = record_date or live_board_state.get("status_date") or datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        self.update_buttons(date_to_use)

    def update_buttons(self, current_date: str):
        today = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        self.btn_today.style = discord.ButtonStyle.primary if current_date == today else discord.ButtonStyle.secondary
        self.btn_indicator.label = f"📅 {current_date}"

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if not await is_ecoda_checker(interaction.user):
            await interaction.response.send_message(
                "❌ **Access Denied:** Only members with the **Checker / Team Leader** role or Server Administrators can use these controls.",
                ephemeral=True,
            )
            return False
        return True

    @discord.ui.button(label="◀️ Prev Day", style=discord.ButtonStyle.secondary, custom_id="live_team_status:prev", row=0)
    async def btn_prev(self, interaction: discord.Interaction, button: discord.ui.Button):
        cur_date_str = live_board_state.get("status_date") or datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        try:
            dt = datetime.strptime(cur_date_str, "%Y-%m-%d") - timedelta(days=1)
            live_board_state["status_date"] = dt.strftime("%Y-%m-%d")
        except Exception:
            live_board_state["status_date"] = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await interaction.response.defer()
        await update_live_team_status_board(interaction.client)

    @discord.ui.button(label="📅 Date", style=discord.ButtonStyle.secondary, disabled=True, custom_id="live_team_status:indicator", row=0)
    async def btn_indicator(self, interaction: discord.Interaction, button: discord.ui.Button):
        pass

    @discord.ui.button(label="Next Day ▶️", style=discord.ButtonStyle.secondary, custom_id="live_team_status:next", row=0)
    async def btn_next(self, interaction: discord.Interaction, button: discord.ui.Button):
        cur_date_str = live_board_state.get("status_date") or datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        try:
            dt = datetime.strptime(cur_date_str, "%Y-%m-%d") + timedelta(days=1)
            live_board_state["status_date"] = dt.strftime("%Y-%m-%d")
        except Exception:
            live_board_state["status_date"] = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await interaction.response.defer()
        await update_live_team_status_board(interaction.client)

    @discord.ui.button(label="📅 Today", style=discord.ButtonStyle.primary, custom_id="live_team_status:today", row=1)
    async def btn_today(self, interaction: discord.Interaction, button: discord.ui.Button):
        live_board_state["status_date"] = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await interaction.response.defer()
        await update_live_team_status_board(interaction.client)

    @discord.ui.button(label="🔄 Refresh", style=discord.ButtonStyle.success, custom_id="live_team_status:refresh", row=1)
    async def btn_refresh(self, interaction: discord.Interaction, button: discord.ui.Button):
        await interaction.response.defer()
        await update_live_team_status_board(interaction.client)


class OnDemandTeamUploadStatusView(discord.ui.View):
    def __init__(self, record_date: str):
        super().__init__(timeout=180)
        self.record_date = record_date
        self.update_buttons()

    def update_buttons(self):
        today = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        self.btn_today.style = discord.ButtonStyle.primary if self.record_date == today else discord.ButtonStyle.secondary
        self.btn_indicator.label = f"📅 {self.record_date}"

    async def interaction_check(self, interaction: discord.Interaction) -> bool:
        if not await is_ecoda_checker(interaction.user):
            await interaction.response.send_message(
                "❌ **Access Denied:** Only members with the **Checker / Team Leader** role or Server Administrators can use these controls.",
                ephemeral=True,
            )
            return False
        return True

    async def _update_view_message(self, interaction: discord.Interaction):
        self.update_buttons()
        status_data = await fetch_team_upload_status(self.record_date)
        embed = build_team_upload_status_embed(interaction.guild, status_data, self.record_date)
        await interaction.response.edit_message(embed=embed, view=self)

    @discord.ui.button(label="◀️ Prev Day", style=discord.ButtonStyle.secondary, row=0)
    async def btn_prev(self, interaction: discord.Interaction, button: discord.ui.Button):
        try:
            dt = datetime.strptime(self.record_date, "%Y-%m-%d") - timedelta(days=1)
            self.record_date = dt.strftime("%Y-%m-%d")
        except Exception:
            self.record_date = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await self._update_view_message(interaction)

    @discord.ui.button(label="📅 Date", style=discord.ButtonStyle.secondary, disabled=True, row=0)
    async def btn_indicator(self, interaction: discord.Interaction, button: discord.ui.Button):
        pass

    @discord.ui.button(label="Next Day ▶️", style=discord.ButtonStyle.secondary, row=0)
    async def btn_next(self, interaction: discord.Interaction, button: discord.ui.Button):
        try:
            dt = datetime.strptime(self.record_date, "%Y-%m-%d") + timedelta(days=1)
            self.record_date = dt.strftime("%Y-%m-%d")
        except Exception:
            self.record_date = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await self._update_view_message(interaction)

    @discord.ui.button(label="📅 Today", style=discord.ButtonStyle.primary, row=1)
    async def btn_today(self, interaction: discord.Interaction, button: discord.ui.Button):
        self.record_date = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
        await self._update_view_message(interaction)

    @discord.ui.button(label="🔄 Refresh", style=discord.ButtonStyle.success, row=1)
    async def btn_refresh(self, interaction: discord.Interaction, button: discord.ui.Button):
        await self._update_view_message(interaction)


async def update_live_team_status_board(bot: commands.Bot):
    """Updates the live team upload status message in the designated channel."""
    channel_id_str = await get_setting("live_leaderboard_channel_id")
    status_msg_id_str = await get_setting("live_status_msg_id")
    if not channel_id_str or not channel_id_str.isdigit():
        return

    channel = bot.get_channel(int(channel_id_str))
    if not channel:
        try:
            channel = await bot.fetch_channel(int(channel_id_str))
        except Exception:
            return

    target_date = live_board_state.get("status_date") or datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
    status_data = await fetch_team_upload_status(target_date)
    embed = build_team_upload_status_embed(channel.guild, status_data, target_date)
    view = LiveTeamUploadStatusView(record_date=target_date)

    if status_msg_id_str and status_msg_id_str.isdigit():
        try:
            msg = await channel.fetch_message(int(status_msg_id_str))
            await msg.edit(content="", embed=embed, view=view)
            return
        except discord.NotFound:
            pass
        except Exception as e:
            print(f"Notice: Failed to update live team status message: {e}")
            return

    # Auto-deploy 3rd message if leaderboard channel is configured but status msg does not exist yet
    try:
        new_msg = await channel.send(embed=embed, view=view)
        await set_setting("live_status_msg_id", str(new_msg.id))
        perms = channel.permissions_for(channel.guild.me)
        if perms.manage_messages:
            try:
                await new_msg.pin(reason="Live ECODA Team Upload Status Board")
            except Exception:
                pass
    except Exception as e:
        print(f"Notice: Failed to deploy live team status message: {e}")


async def update_live_leaderboard_messages(bot: commands.Bot):
    """Refreshes Voice, ECODA, and Team Upload Status live boards in the designated channel."""
    await update_live_voice_leaderboard(bot)
    await update_live_ecoda_leaderboard(bot)
    await update_live_team_status_board(bot)

